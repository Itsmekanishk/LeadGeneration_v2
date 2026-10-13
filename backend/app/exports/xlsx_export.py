"""XLSX export builder (openpyxl)."""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_COLUMNS = [
    "Business Name", "Category", "Address", "Phone", "Email", "Website",
    "Rating", "Review Count", "Email Confidence", "Phone Source",
    "Website Source", "All Emails", "All Phones",
    "Search Query", "Search Location", "Collected At",
]
_WIDTHS = [34, 20, 46, 18, 32, 36, 8, 13, 16, 14, 22, 44, 26, 34, 26, 20]

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

# Signed plain numbers, e.g. an E.164 phone "+919876543210". Kept verbatim so
# the exported phone number stays importable into a CRM.
_NUMERIC_LIKE_RE = re.compile(r"^[+\-]?[\d\s().\-]*\d$")


def _safe_cell(value):
    """Neutralize spreadsheet formula injection.

    openpyxl writes string values starting with ``=`` as *live formulas*, so
    business-derived text like ``=HYPERLINK(...)`` or ``=cmd|...`` would
    execute on open. Prefixing an apostrophe forces text interpretation.
    Plain numbers and phone numbers are exempt so their value is preserved.
    """
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        if _NUMERIC_LIKE_RE.match(value):
            return value
        return f"'{value}"
    return value


def build_xlsx(out_path: Path, rows: list[dict]) -> None:
    """Write a styled workbook: bold header, freeze panes, autofilter."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Businesses"

    header_fill = PatternFill("solid", fgColor="1D4ED8")
    header_font = Font(color="FFFFFF", bold=True, size=10)

    for col, name in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = _WIDTHS[col - 1]

    for r, row in enumerate(rows, start=2):
        for c, name in enumerate(_COLUMNS, start=1):
            value = row.get(name)
            if name == "Collected At" and isinstance(value, datetime):
                value = value.strftime("%Y-%m-%d %H:%M:%S")
            elif name == "Rating" and isinstance(value, float):
                value = round(value, 1)
            cell = ws.cell(row=r, column=c, value=_safe_cell(value))
            if name == "Website" and isinstance(value, str) and value.startswith("http"):
                cell.hyperlink = value
                cell.style = "Hyperlink"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_COLUMNS))}{max(len(rows) + 1, 1)}"
    ws.row_dimensions[1].height = 22
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
